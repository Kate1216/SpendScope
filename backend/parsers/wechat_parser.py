import csv
from hashlib import sha1
from io import BytesIO, StringIO
import re
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - handled by caller environment
    load_workbook = None


TIME_FIELDS = ["交易时间", "支付时间", "创建时间", "记账日期", "交易日期", "时间"]
MERCHANT_FIELDS = ["交易对方", "商户", "商家", "收/付款方", "交易对象", "对方"]
DESCRIPTION_FIELDS = ["商品", "商品说明", "交易说明", "备注", "摘要", "说明"]
DIRECTION_FIELDS = ["收/支", "收支", "收/支状态"]
AMOUNT_FIELDS = ["金额", "交易金额", "金额(元)", "交易金额(元)"]
PAYMENT_FIELDS = ["支付方式", "收/付款方式", "付款方式", "收款方式"]
TRANSACTION_TYPE_FIELDS = ["交易类型", "业务类型", "账务类型"]
STATUS_FIELDS = ["当前状态", "交易状态"]


def parse_wechat_file(file_bytes: bytes, filename: str = "") -> list[dict]:
    rows = _read_excel_rows(file_bytes, filename) if _is_excel(filename) else _read_csv_rows(file_bytes)
    if not rows:
        return []

    header_index = _find_header_index(rows)
    if header_index < 0:
        return []

    headers = [_clean_cell(cell) or f"列{index + 1}" for index, cell in enumerate(rows[header_index])]
    records = [_row_to_record(headers, row) for row in rows[header_index + 1 :]]

    transactions = []
    for record in records:
        transaction = _record_to_transaction(record)
        if transaction:
            transactions.append(transaction)

    return transactions


def _read_csv_rows(file_bytes: bytes) -> list[list[str]]:
    text = None
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            text = file_bytes.decode(encoding)
            break
        except UnicodeDecodeError:
            continue

    if text is None:
        text = file_bytes.decode("utf-8", errors="ignore")

    lines = [line for line in text.splitlines() if line.strip()]
    if not lines:
        return []

    delimiter = _detect_delimiter(lines)
    reader = csv.reader(StringIO("\n".join(lines)), delimiter=delimiter)
    return [[_clean_cell(cell) for cell in row] for row in reader]


def _read_excel_rows(file_bytes: bytes, filename: str) -> list[list[str]]:
    if load_workbook is None:
        return []

    workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        rows = []
        for sheet in workbook.worksheets:
            sheet_rows = [[_clean_cell(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
            if _find_header_index(sheet_rows) >= 0:
                return sheet_rows
            rows.extend(sheet_rows)
        return rows
    finally:
        workbook.close()


def _is_excel(filename: str) -> bool:
    return str(filename or "").lower().endswith(".xlsx")


def _detect_delimiter(lines: list[str]) -> str:
    sample = "\n".join(lines[:10])
    try:
        return csv.Sniffer().sniff(sample, delimiters=",\t;").delimiter
    except csv.Error:
        first = lines[0] if lines else ""
        if "\t" in first:
            return "\t"
        if ";" in first:
            return ";"
        return ","


def _find_header_index(rows: list[list[Any]]) -> int:
    for index, row in enumerate(rows):
        cells = [_clean_cell(cell) for cell in row]
        text = "|".join(cells)
        has_time = any(field in text for field in TIME_FIELDS)
        has_amount_or_direction = any(field in text for field in AMOUNT_FIELDS + DIRECTION_FIELDS + ["收入", "支出"])
        if has_time and has_amount_or_direction:
            return index
    return -1


def _row_to_record(headers: list[str], row: list[Any]) -> dict:
    return {
        header: _clean_cell(row[index] if index < len(row) else "")
        for index, header in enumerate(headers)
    }


def _record_to_transaction(record: dict) -> dict | None:
    time = _pick(record, TIME_FIELDS)
    amount = _parse_amount(_pick(record, AMOUNT_FIELDS))
    if amount is None or amount == 0:
        return None

    merchant = _pick(record, MERCHANT_FIELDS) or _pick(record, DESCRIPTION_FIELDS) or "未知交易对象"
    description = _pick(record, DESCRIPTION_FIELDS) or merchant
    platform = _pick(record, PAYMENT_FIELDS)
    transaction_type = _pick(record, TRANSACTION_TYPE_FIELDS)
    status = _pick(record, STATUS_FIELDS)
    type_result = _normalize_wechat_type(_pick(record, DIRECTION_FIELDS), transaction_type, status)
    tx_type = type_result.get("type")
    if not tx_type:
        return None

    transaction = {
        "time": time,
        "sourcePlatform": "微信",
        "platform": platform,
        "merchant": merchant,
        "description": description,
        "transactionType": transaction_type,
        "type": tx_type,
        "category": type_result.get("category") or _default_category(tx_type),
        "amount": abs(amount),
    }
    if type_result.get("excludeReason"):
        transaction["excludeReason"] = type_result["excludeReason"]
    if type_result.get("needsReview"):
        transaction["needsReview"] = True

    transaction["id"] = _build_transaction_id(transaction)
    return transaction


def _normalize_wechat_type(direction: str, trade_type: str, status: str) -> dict:
    direction = _clean_cell(direction)
    trade_type = _clean_cell(trade_type)
    status = _clean_cell(status)
    fallback_text = f"{trade_type} {status}"

    if re.search(r"不计\s*收支|不计入收支|不计", direction):
        return {"type": "排除", "category": "排除", "excludeReason": "不计收支"}
    if "支出" in direction:
        return {"type": "支出"}
    if "收入" in direction:
        return {"type": "收入"}

    if re.search(r"不计\s*收支|不计入收支|不计", fallback_text):
        return {"type": "排除", "category": "排除", "excludeReason": "不计收支"}
    if re.search(r"退款|退货|售后退款", fallback_text):
        return {"type": "退款", "category": "退款/抵扣"}
    if re.search(r"转账收款|收款|红包收入|收红包|已收钱|入账", fallback_text):
        return {"type": "收入"}
    if re.search(r"商户消费|扫二维码付款|二维码付款|微信支付|群收款付款|付款给商户|消费", fallback_text):
        return {"type": "支出"}
    if re.search(r"^转账$|转账", trade_type):
        return {"type": "排除", "category": "排除", "excludeReason": "转账待确认", "needsReview": True}

    return {}


def _default_category(tx_type: str) -> str:
    if tx_type == "支出":
        return "待确认"
    if tx_type == "收入":
        return "其他"
    if tx_type == "退款":
        return "退款/抵扣"
    if tx_type == "排除":
        return "排除"
    return "待确认"


def _pick(record: dict, aliases: list[str]) -> str:
    for alias in aliases:
        for key, value in record.items():
            if _same_header(key, alias):
                return _clean_cell(value)
    return ""


def _same_header(left: str, right: str) -> bool:
    return _normalize_header(left) == _normalize_header(right)


def _normalize_header(value: str) -> str:
    return re.sub(r"[\s：:()（）/\\_-]+", "", str(value or "")).lower()


def _clean_cell(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\ufeff", "").replace("\n", " ")).strip()


def _parse_amount(value: str) -> float | None:
    text = _clean_cell(value).replace(",", "")
    match = re.search(r"[-+]?\d+(?:\.\d+)?", text)
    if not match:
        return None
    try:
        return float(match.group())
    except ValueError:
        return None


def _build_transaction_id(transaction: dict) -> str:
    raw = "|".join(
        [
            transaction.get("sourcePlatform") or "",
            transaction.get("time") or "",
            transaction.get("merchant") or "",
            transaction.get("description") or "",
            str(transaction.get("amount") or 0),
            transaction.get("type") or "",
        ]
    )
    return f"backend-wechat-{sha1(raw.encode('utf-8')).hexdigest()[:16]}"
